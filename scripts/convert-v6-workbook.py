#!/usr/bin/env python3
"""Convert Ingredients_Sauces_Dressings_v6.xlsx into the two committed JSON documents.

    python scripts/convert-v6-workbook.py <path-to-xlsx>   # writes both JSONs
    python scripts/convert-v6-workbook.py --self-test      # only the built-in checks

Outputs (paths relative to the repo root this script lives under):
  apps/api/app-modules/ingredients/database/data/platform-ingredients.json
      Sheets "1. Ingredients" and "6.Packaging" -> the platform ingredient
      library (categories, ingredients, aliases) in the exact document shape
      IngredientMasterSeeder reads: 337 rows, ING-001..306 and PKG-001..031.

      Packaging shares the ingredient table and is told apart by its category:
      every row from sheet 6 is filed under `packaging-disposables`, the
      ingredient list excludes that branch and the packaging list asks for it
      by name. It was briefly its own table, and the argument for that was a
      bug rather than a difference — a category filter that silently degraded
      to no filter — which is now closed on both sides at the repository.

      Sheet 6 carries no cost column, so the packaging rows' prices and
      capacities come from the curated overlay beside this script's output
      (`packaging-overlay.json`), merged in by source_ref. Those figures are
      indicative foodservice wholesale, not transcribed — the overlay states
      its own provenance and the seed document carries it forward.
  apps/api/app-modules/kitchens/database/data/v6-catalogue.json
      Sheets 2-5 (sauces, dressings, meals, resale products) -> the org
      catalogue the kitchen:import-v6 command reads. Prices sit in their own
      per-row section so they could be split into a separate file without
      touching anything else.

The workbook has known data defects; every repair is implemented here, once,
so the committed JSON is clean. Each repaired row carries a machine-readable
flag in its "flags" list, and the run prints a repair summary. The importer
and seeder stay dumb. --self-test builds a synthetic in-memory workbook that
exercises every repair rule plus plain parsing and asserts the normalized
output; it also runs automatically before every real conversion.
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
PLATFORM_OUT = REPO_ROOT / "apps/api/app-modules/ingredients/database/data/platform-ingredients.json"
CATALOGUE_OUT = REPO_ROOT / "apps/api/app-modules/kitchens/database/data/v6-catalogue.json"
# Curated packaging prices and capacities. Not workbook-derived — sheet 6 has no cost column —
# so it is an input to the conversion rather than an output of it.
PACKAGING_OVERLAY = REPO_ROOT / "apps/api/app-modules/ingredients/database/data/packaging-overlay.json"

# ---------------------------------------------------------------- vocabulary

# Ported from app-modules/kitchens/src/Import/Runtime/AllergenClassMap.php —
# a closed table, no fuzzy matching. The v6 sheets additionally use the
# singular "Tree nut" and sometimes leave the *(US-only coconut) marker only
# in the Allergen Source column, so resolution looks at both cells.
ALLERGEN_CODES = {
    "milk": "milk",
    "eggs": "egg",
    "egg": "egg",
    "fish": "fish",
    "crustaceans": "crustaceans",
    "molluscs": "mollusc",
    "tree nuts": "tree_nut",
    "tree nut": "tree_nut",
    "peanuts": "peanut",
    "sesame": "sesame",
    "cereals/gluten": "gluten",
    "cereals / gluten": "gluten",
    "gluten": "gluten",
    "soybeans": "soy",
    "soya": "soy",
    "mustard": "mustard",
    "celery": "celery",
    "sulphites": "sulphites",
    "lupin": "lupin",
}

# Workbook unit spellings -> measurement_units codes (kitchens UnitMap dialect).
UNIT_CODES = {
    "KG": "kg", "G": "g", "GR": "g", "LTR": "l", "L": "l", "ML": "ml",
    "PIECE": "piece", "PC": "piece", "BAG": "bag", "CAN": "can",
    "GAL": "gallon", "GALLON": "gallon", "PACK": "pack", "BTL": "bottle",
    "BOTTLE": "bottle", "BUNCH": "bunch",
}

# Sheet13 of the workbook: purchase TYPE -> usage UNIT, used only to default
# a blank Unit cell from its Type.
TYPE_TO_USAGE = {
    "BAG": "PIECE", "BTL": "LTR", "CAN": "PIECE", "GAL": "LTR",
    "KG": "KG", "LTR": "LTR", "PACK": "PIECE", "PIECE": "PIECE",
}

# Published Category wording -> product_categories codes.
PUBLISHED_CATEGORY_CODES = {
    "sauce & marinade": "sauce",
    "dressing": "dressing",
    "meal": "meal",
    "frozen": "frozen",
    "beverage": "beverage",
    "bread": "bread",
}

# Case/plural variants folded onto one taxonomy node. Keys and values are the
# verbatim sheet wording; resolution happens before slugging so "Vegetables"
# and "Vegetable" become one platform category.
CATEGORY_SYNONYMS = {
    "vegetables": "Vegetable",
    "mushroom": "Mushrooms",
    "cake & cream": "Cakes & Cream",
    "leafy green": "Leafy Green",
    "tree nut": "Tree Nut",
    "dressings": "Dressings",
}


def slugify(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    assert slug, f"unsluggable value: {value!r}"
    return slug


def clean(cell) -> str:
    if cell is None:
        return ""
    return re.sub(r"\s+", " ", str(cell)).strip()


def canonical_category(name: str) -> str:
    return CATEGORY_SYNONYMS.get(name.lower(), name)


def unit_code(cell: str, *, default: str | None = None) -> str | None:
    token = clean(cell).upper().rstrip(".")
    if not token or token == "-":
        return default
    code = UNIT_CODES.get(token)
    assert code is not None, f"unknown unit spelling: {cell!r}"
    return code


ALLERGEN_TOKEN = re.compile(
    r"^\s*(milk|eggs?|fish|crustaceans|molluscs|tree nuts?|peanuts|sesame|"
    r"cereals\s*/\s*gluten|gluten|soybeans|soya|mustard|celery|sulphites|lupin|none)\s*[~*]?\s*([-—+,;]|$)",
    re.IGNORECASE,
)


def looks_like_allergen_source(text: str) -> bool:
    """Whether a cell reads as allergen-source shorthand ("Milk - butter; ...")
    or an explicit dash/None, as opposed to a composition list."""
    t = clean(text)
    if t in {"", "-", "—"} or t.lower() == "none":
        return True
    return bool(ALLERGEN_TOKEN.match(t))


def resolve_allergens(class_cell: str, source_cell: str):
    """Allergen Class + Allergen Source cells -> (mappings, evidence, flags).

    Class is a comma list of workbook wordings; markers `*` (US-only coconut)
    and `~` (sulphites: verify per supplier) may sit in either cell.
    """
    flags: list[str] = []
    text = clean(class_cell)
    evidence_text = clean(source_cell)
    if evidence_text in {"-", "—"}:
        evidence_text = ""

    if not text or text.lower() == "none":
        return [], evidence_text or None, flags

    # Per-code evidence where the source text splits cleanly ("Eggs - yolk; ...").
    evidence_by_code: dict[str, str] = {}
    for part in evidence_text.split(";"):
        part = part.strip()
        head = re.split(r"\s[-—]\s|\s-\s?", part, maxsplit=1)[0] if part else ""
        for token in re.split(r"\s*[+,]\s*", head):
            key = re.sub(r"\s+", " ", token.replace("*", "").replace("~", "").strip().lower())
            code = ALLERGEN_CODES.get(key)
            if code and part:
                evidence_by_code[code] = part

    mappings = []
    seen = set()
    for token in text.split(","):
        raw = token.strip()
        if not raw:
            continue
        us_only = "*" in raw
        verify = "~" in raw
        key = re.sub(r"\s+", " ", raw.replace("*", "").replace("~", "").strip().lower())
        if key in {"", "none", "-", "—"}:
            continue
        code = ALLERGEN_CODES.get(key)
        if code is None:
            flags.append(f"allergen_unmapped:{raw}")
            continue
        # The v6 sheets sometimes leave the coconut marker only in the source
        # column ("Tree nuts* - ..."), with a bare "Tree nut" class.
        if code == "tree_nut" and "tree nuts*" in evidence_text.lower():
            us_only = True
        if code == "sulphites" and "~" in evidence_text:
            verify = True
        if code in seen:
            continue
        seen.add(code)
        mapping = {"allergen_code": code}
        if us_only:
            mapping["market_scope"] = "us_only"
        if verify:
            # `~` is the workbook's "possible — verify per supplier": a
            # possibility, not a determination, on both axes.
            mapping["containment"] = "may_contain"
            mapping["verification_status"] = "requires_supplier_confirmation"
        per_code = evidence_by_code.get(code)
        if per_code:
            mapping["evidence"] = per_code
        mappings.append(mapping)

    if not mappings and not flags:
        flags.append(f"allergen_unmapped:{text}")
    return mappings, evidence_text or None, flags


def money_minor(cell) -> int | None:
    """A price cell -> integer US-cent minor units. Accepts numbers and
    '$5.00'-style strings; blank/dash -> None."""
    if cell is None:
        return None
    if isinstance(cell, (int, float)):
        return round(float(cell) * 100)
    text = clean(cell).replace("$", "").replace(",", "")
    if text in {"", "-", "—"}:
        return None
    return round(float(text) * 100)


def number_or_none(cell) -> float | None:
    if cell is None or clean(cell) in {"", "-", "—"}:
        return None
    return float(cell)


# ---------------------------------------------------------------- sheet access

def sheet_rows(ws, id_prefix: str):
    """Data rows of a v6 sheet: header on row 3, data from row 4, keyed rows
    only (ID startswith the prefix)."""
    rows = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        row_id = clean(row[0]) if row else ""
        if row_id.startswith(id_prefix):
            rows.append(row)
    return rows


class SlugBook:
    """Per-namespace slug allocation; collisions get a deterministic suffix."""

    def __init__(self):
        self.taken: dict[str, set[str]] = {}

    def claim(self, namespace: str, name: str, *, suffix: str | None = None) -> tuple[str, bool]:
        book = self.taken.setdefault(namespace, set())
        base = slugify(name)
        slug = base
        collided = False
        if slug in book and suffix:
            slug, collided = f"{base}-{suffix}", True
        n = 2
        while slug in book:
            slug, collided = f"{base}-{n}", True
            n += 1
        book.add(slug)
        return slug, collided


class Taxonomy:
    """The platform category tree, unioned over every sheet's Category /
    Sub-Category pair in first-appearance order."""

    def __init__(self):
        self.parents: dict[str, dict] = {}
        self.children: dict[str, dict] = {}
        self.order = 0

    def add(self, category: str, subcategory: str | None):
        category = canonical_category(clean(category))
        if not category:
            return None, None
        parent_code = slugify(category)
        if parent_code not in self.parents:
            self.order += 1
            self.parents[parent_code] = {
                "code": parent_code, "name_en": category,
                "parent_code": None, "display_order": self.order,
            }
        sub_code = None
        subcategory = canonical_category(clean(subcategory or ""))
        if subcategory:
            sub_code = f"{parent_code}-{slugify(subcategory)}"
            if sub_code not in self.children:
                self.order += 1
                self.children[sub_code] = {
                    "code": sub_code, "name_en": subcategory,
                    "parent_code": parent_code, "display_order": self.order,
                }
        return parent_code, sub_code

    def rows(self):
        return list(self.parents.values()) + list(self.children.values())


# ---------------------------------------------------------------- converters

def convert_ingredients(ws, taxonomy: Taxonomy, slugs: SlugBook, repairs: list[str]):
    """Sheet '1. Ingredients' -> platform ingredient rows."""
    out = []
    for row in sheet_rows(ws, "ING-"):
        (rid, item, _published, category, subcategory, _ing, _prod, kind, status,
         allergen_class, composition, allergen_source, type_cell, unit_cell,
         items_per_unit, _price, notes) = (list(row) + [None] * 17)[:17]
        rid, item = clean(rid), clean(item)
        flags = []

        category_code, subcategory_code = taxonomy.add(category, subcategory)
        status_value = "active" if clean(status).lower() == "active" else "inactive"
        if status_value == "inactive":
            flags.append("status_blank_inactive")
            repairs.append(f"{rid}: blank Status -> inactive")

        usage = unit_code(unit_cell) or unit_code(TYPE_TO_USAGE.get(clean(type_cell).upper(), ""))
        if usage is None:
            usage = "kg"
            flags.append("unit_defaulted_kg")
        purchase = unit_code(type_cell, default=usage)

        mappings, evidence, allergen_flags = resolve_allergens(allergen_class or "", allergen_source or "")
        flags += allergen_flags
        for mapping in mappings:
            if evidence and "evidence" not in mapping:
                mapping["evidence"] = evidence

        slug, collided = slugs.claim("platform-ingredient", item)
        if collided:
            flags.append("slug_collision_suffixed")

        # Every row of the ingredient master is a supplier good. The sheet's
        # Kind column disagrees exactly once (ING-278 Caramelised Onions,
        # which is also PRD-005 on the production sheet), and that is the
        # catalogue row's fact, not this one's: whether something is made
        # in-house is a property of the recipe that outputs it, never of the
        # ingredient. The override is flagged rather than silent.
        if clean(kind).lower() not in {"", "supplier", "suplier"}:
            flags.append("kind_forced_supplier")
            repairs.append(f"{rid}: Kind {clean(kind)!r} -> supplier")

        entry = {
            "source_ref": rid,
            "slug": slug,
            "name_en": item,
            "category_code": category_code,
            "subcategory_code": subcategory_code,
            "default_unit_code": usage,
            "purchase_unit_code": purchase,
            "status": status_value,
            "kind": "supplier",
            "allergens": mappings,
        }
        if clean(composition):
            entry["composition"] = clean(composition)
        ipu = number_or_none(items_per_unit)
        if ipu is not None:
            entry["items_per_unit"] = ipu
        if clean(notes):
            entry["notes"] = clean(notes)
        if flags:
            entry["flags"] = sorted(set(flags))
        out.append(entry)
    return out


def convert_packaging(ws, taxonomy: Taxonomy, slugs: SlugBook, overlay: dict):
    """Sheet '6.Packaging' -> platform ingredient rows.

    Non-food supplier goods: no allergens, no nutrition, no yield. They are
    ingredients in storage and in shape, and packaging by category.

    `overlay` is the curated price/capacity document keyed by source_ref. The
    sheet has no cost column at all, so without it every box seeds unpriced -
    and a recipe's technical sheet withholds its total entirely while any one
    packaging line has no price, so one unpriced box takes the bottom line off
    every recipe that ships in it.
    """
    out = []
    for row in sheet_rows(ws, "PKG-"):
        (rid, item, category, subcategory, _kind, status, composition,
         type_cell, unit_cell, items_per_unit, waste) = (list(row) + [None] * 11)[:11]
        rid, item = clean(rid), clean(item)
        category_code, subcategory_code = taxonomy.add(category, subcategory)
        usage = unit_code(unit_cell) or "piece"
        slug, _ = slugs.claim("platform-ingredient", item)
        entry = {
            "source_ref": rid,
            "slug": slug,
            "name_en": item,
            "category_code": category_code,
            "subcategory_code": subcategory_code,
            "default_unit_code": usage,
            "purchase_unit_code": unit_code(type_cell, default=usage),
            "status": "active" if clean(status).lower() == "active" else "inactive",
            "kind": "supplier",
            "allergens": [],
        }
        if clean(composition):
            entry["composition"] = clean(composition)
        ipu = number_or_none(items_per_unit)
        if ipu is not None:
            entry["items_per_unit"] = ipu
        pct = number_or_none(waste)
        if pct is not None:
            entry["waste_percent"] = pct

        # The curated half. Only keys the overlay actually states are copied, so a row it does
        # not mention seeds exactly as the sheet describes it rather than with invented nulls.
        #
        # `items_per_unit` is in this list and is the reason the list is not just prices: sheet 6
        # leaves the column blank on every row, and it is the divisor between a pack price and a
        # per-piece cost. Without it a case of 500 lids at 40.00 costs a recipe 40.00 a lid. The
        # sheet still wins where it states one — the overlay only fills a gap.
        for field in ("items_per_unit", "purchase_price_amount", "purchase_price_currency",
                      "capacity_quantity", "capacity_unit_code"):
            if entry.get(field) is not None:
                continue
            value = overlay.get(rid, {}).get(field)
            if value is not None:
                entry[field] = value

        out.append(entry)
    return out


# Column layout shared by sheets 2 (sauce), 3 (dressing) and 4 (meal); sheet 3
# titles its unit column "Usage Unit", same position. Sheet 4 has no
# Pieces/Pack column.
PRODUCTION_COLUMNS = [
    "id", "name", "published", "category", "subcategory", "ingredient",
    "product", "kind", "status", "source", "composition", "allergen_class",
    "allergen_source", "type", "unit", "items_per_unit", "b2b_weight",
    "b2b_price", "b2c_weight", "b2c_price",
]


def convert_production_sheet(ws, id_prefix: str, item_type: str, taxonomy: Taxonomy,
                             slugs: SlugBook, repairs: list[str]):
    has_pieces = item_type in {"sauce", "dressing"}
    out = []
    for row in sheet_rows(ws, id_prefix):
        cells = list(row) + [None] * 24
        record = dict(zip(PRODUCTION_COLUMNS, cells))
        tail_start = len(PRODUCTION_COLUMNS)
        if has_pieces:
            record["pieces_per_pack"] = cells[tail_start]
            tail_start += 1
        # Notes plus any overflow cells beyond the header (the sheets carry
        # spill-over text in unnamed columns).
        notes = " | ".join(clean(c) for c in cells[tail_start:] if clean(c))

        rid, name = clean(record["id"]), clean(record["name"])
        flags = []

        composition = clean(record["composition"])
        allergen_class = clean(record["allergen_class"])
        # PRD-016 pattern: composition text sits in the Allergen Class column
        # and "None"/nothing in the composition column.
        if looks_like_allergen_source(composition) and not looks_like_allergen_source(allergen_class):
            composition, allergen_class = allergen_class, composition
            flags.append("columns_swapped_repaired")
            repairs.append(f"{rid}: composition/allergen columns un-swapped")

        ingredient_cell = clean(record["ingredient"]).lower()
        product_cell = clean(record["product"]).lower()
        is_ingredient = ingredient_cell == "yes"
        is_product = product_cell == "yes"
        if not ingredient_cell and not product_cell:
            # Four meal rows leave both blank (one of them priced): sellable is
            # the only consistent reading; the doubt is preserved as a flag.
            is_product, is_ingredient = True, False
            flags.append("source_blank_role_flags")
            repairs.append(f"{rid}: blank Ingredient?/Product? -> product only")

        usage = unit_code(record["unit"]) or unit_code(TYPE_TO_USAGE.get(clean(record["type"]).upper(), ""))
        if usage is None:
            usage = "kg"
            flags.append("unit_defaulted_kg")
        mappings, evidence, allergen_flags = resolve_allergens(allergen_class, clean(record["allergen_source"]))
        flags += allergen_flags

        category_code, subcategory_code = taxonomy.add(record["category"], record["subcategory"])
        slug, collided = slugs.claim("catalogue-item", name, suffix=item_type)
        if collided:
            flags.append("slug_collision_suffixed")
            repairs.append(f"{rid}: slug collision -> {slug}")

        source = clean(record["source"])
        if source.lower() == "recipe library":
            flags.append("recipe_library_unlinked")

        prices = {}
        b2b, b2c = money_minor(record["b2b_price"]), money_minor(record["b2c_price"])
        if b2b is not None:
            prices["b2b"] = {"weight_kg": number_or_none(record["b2b_weight"]), "price_minor": b2b, "currency": "USD"}
        if b2c is not None:
            prices["b2c"] = {"weight_kg": number_or_none(record["b2c_weight"]), "price_minor": b2c, "currency": "USD"}

        entry = {
            "source_ref": rid,
            "sheet_item_type": item_type,
            "name_en": name,
            "slug": slug,
            "published_category_code": PUBLISHED_CATEGORY_CODES.get(clean(record["published"]).lower()),
            "kitchen_category": clean(record["category"]) or None,
            "kitchen_subcategory": clean(record["subcategory"]) or None,
            "is_ingredient": is_ingredient,
            "is_product": is_product,
            "production_mode": "production",
            "status": "active" if clean(record["status"]).lower() == "active" else "draft",
            "source": source or None,
            "composition": composition or None,
            "category_code": category_code,
            "subcategory_code": subcategory_code,
            "usage_unit_code": usage,
            "purchase_unit_code": unit_code(record["type"], default=usage),
            "items_per_unit": number_or_none(record["items_per_unit"]),
            "pieces_per_pack": number_or_none(record.get("pieces_per_pack")),
            "allergens": mappings,
            "allergen_evidence": evidence,
            "prices": prices,
            "notes": notes or None,
            "flags": sorted(set(flags)),
        }
        out.append(entry)
    return out


# Sheet 5 header: ID | Item | Published | Category | Sub-Category |
# Ingredient? | Kind | Status | Recipe (Made From) | Allergen Class |
# Allergen Source | Type | Unit | ... (no Product? column, no prices filled).
def convert_products_sheet(ws, taxonomy: Taxonomy, slugs: SlugBook, repairs: list[str]):
    out = []
    for row in sheet_rows(ws, "RSL-"):
        cells = [clean(c) for c in (list(row) + [None] * 20)[:20]]
        (rid, name, published, category, subcategory, ingredient_cell, kind,
         status, recipe_cell, class_cell, source_cell, type_cell, unit_cell,
         *_rest) = cells
        flags = []

        # RSL-055 pattern: the row follows the ingredients-sheet shape and is
        # shifted one column right from this sheet's header (an extra
        # Product? cell after Ingredient?).
        if kind.lower() == "no" and status.lower() in {"supplier", "suplier", "production"}:
            cells.pop(6)
            cells.append("")
            (rid, name, published, category, subcategory, ingredient_cell, kind,
             status, recipe_cell, class_cell, source_cell, type_cell, unit_cell,
             *_rest) = cells
            flags.append("row_shift_repaired")
            repairs.append(f"{rid}: ingredients-shaped row unshifted")

        if kind.lower() == "suplier":
            flags.append("kind_typo_repaired")

        # RSL-045 pattern: "Recipe Library" sits in the composition column and
        # the composition in the Allergen Class column.
        source = None
        if recipe_cell.lower() == "recipe library":
            source = "Recipe Library"
            recipe_cell = ""
            flags.append("source_misplaced_repaired")
            repairs.append(f"{rid}: 'Recipe Library' moved to Source")
        if not looks_like_allergen_source(class_cell):
            if looks_like_allergen_source(recipe_cell):
                recipe_cell, class_cell = class_cell, recipe_cell
            else:
                recipe_cell, class_cell = class_cell, "None"
            flags.append("columns_swapped_repaired")
            repairs.append(f"{rid}: composition found in Allergen Class column")

        # Most RSL rows carry the composition in the Allergen Source column
        # and the allergen shorthand (or a dash) in the Recipe column.
        composition, evidence_cell = recipe_cell, source_cell
        if looks_like_allergen_source(recipe_cell) and not looks_like_allergen_source(source_cell):
            composition, evidence_cell = source_cell, recipe_cell
            flags.append("columns_swapped_repaired")
            repairs.append(f"{rid}: composition/allergen-source un-swapped")

        usage = unit_code(unit_cell) or unit_code(TYPE_TO_USAGE.get(type_cell.upper(), ""))
        if usage is None:
            usage = "piece"  # resale goods sell as units; the system's own product fallback
            flags.append("unit_defaulted_piece")
        mappings, evidence, allergen_flags = resolve_allergens(class_cell, evidence_cell)
        flags += allergen_flags

        category_code, subcategory_code = taxonomy.add(category, subcategory)
        slug, collided = slugs.claim("catalogue-item", name, suffix="product")
        if collided:
            flags.append("slug_collision_suffixed")
            repairs.append(f"{rid}: slug collision -> {slug}")
        if source and source.lower() == "recipe library":
            flags.append("recipe_library_unlinked")
        if not clean(published):
            flags.append("published_category_missing")

        out.append({
            "source_ref": rid,
            "sheet_item_type": "product",
            "name_en": name,
            "slug": slug,
            "published_category_code": PUBLISHED_CATEGORY_CODES.get(clean(published).lower()),
            "kitchen_category": category or None,
            "kitchen_subcategory": subcategory or None,
            "is_ingredient": ingredient_cell.lower() == "yes",
            "is_product": True,
            "production_mode": "supplier",
            "status": "active" if status.lower() == "active" else "draft",
            "source": source,
            "composition": composition or None,
            "category_code": category_code,
            "subcategory_code": subcategory_code,
            "usage_unit_code": usage,
            "purchase_unit_code": unit_code(type_cell, default=usage),
            "items_per_unit": None,
            "pieces_per_pack": None,
            "allergens": mappings,
            "allergen_evidence": evidence,
            "prices": {},
            "notes": None,
            "flags": sorted(set(flags)),
        })
    return out


# ---------------------------------------------------------------- driver

def worksheet(workbook, *names):
    """Resolve a worksheet by any of the titles the export has used for it.

    Exports rename tabs between revisions — sheet 4 has shipped as both
    "4. Meals" and "4. Production", sheet 5 as both "5. Products" and
    "5. Resale" — while the column layout and the row identifier prefix stay
    put. Rows are classified by their own `PRD-`/`RSL-` identifier anyway, so
    a tab rename must not be the thing that breaks a conversion. Matching
    ignores case and whitespace; an unknown title still fails loudly, naming
    what the workbook actually contains.
    """
    available = {str(t).strip().lower(): t for t in workbook.sheetnames}
    for name in names:
        title = available.get(name.strip().lower())
        if title is not None:
            return workbook[title]
    raise KeyError(f"no sheet named any of {list(names)}; workbook has {workbook.sheetnames}")


def convert(workbook, packaging_overlay=None):
    taxonomy = Taxonomy()
    slugs = SlugBook()
    repairs: list[str] = []

    platform = convert_ingredients(worksheet(workbook, "1. Ingredients"), taxonomy, slugs, repairs)
    packaging = convert_packaging(
        worksheet(workbook, "6.Packaging"), taxonomy, slugs, packaging_overlay or {},
    )
    catalogue = (
        convert_production_sheet(worksheet(workbook, "2. Sauce & Marination"), "SAC-", "sauce", taxonomy, slugs, repairs)
        + convert_production_sheet(worksheet(workbook, "3. Dressing"), "DRS-", "dressing", taxonomy, slugs, repairs)
        + convert_production_sheet(worksheet(workbook, "4. Meals", "4. Production"), "PRD-", "meal", taxonomy, slugs, repairs)
        + convert_products_sheet(worksheet(workbook, "5. Products", "5. Resale"), taxonomy, slugs, repairs)
    )

    platform_doc = {
        "categories": taxonomy.rows(),
        "ingredients": platform + packaging,
        "aliases": [],
    }
    catalogue_doc = {
        "source_workbook": "Ingredients_Sauces_Dressings_v6.xlsx",
        "source_system": "healthy360_workbook_v6",
        "items": catalogue,
    }
    return platform_doc, catalogue_doc, repairs


# ---------------------------------------------------------------- self-test

def build_synthetic_workbook():
    """A tiny in-memory workbook exercising every repair rule plus plain rows."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def fill(title, header, rows):
        ws = wb.create_sheet(title)
        ws.append([title])
        ws.append([])
        ws.append(header)
        for r in rows:
            ws.append(r)
        return ws

    fill("1. Ingredients",
         ["ID", "Item", "Published Category", "Category", "Sub-Category", "Ingredient?", "Product?",
          "Kind", "Status", "Allergen Class", "Composition (Made From)", "Allergen Source",
          "Type", "Unit", "Items per Unit", "Price", "Notes"],
         [
             ["ING-001", "Test Flour", "Baking", "Baking", "Wheat / gluten", "Yes", "No", "Supplier",
              "Active", "Cereals/Gluten", "Milled wheat", "Cereals/Gluten - milled wheat", "KG", "KG", 1, None, "note"],
             # blank status -> inactive; blank units -> kg default
             ["ING-002", "Test Vinegar", "Cond", "Cond", "Condiment", "Yes", "No", "Supplier",
              None, "Sulphites~", None, "Sulphites~ - vinegar", None, None, None, None, None],
             # coconut: singular class, marker only in the source column
             ["ING-003", "Test Coconut", "Fruit", "Fruit", "Dried Fruit", "Yes", "No", "Supplier",
              "Active", "Tree nut", None, "Tree nuts* - coconut", "PACK", None, None, None, None],
         ])
    fill("2. Sauce & Marination",
         ["ID", "Production Item", "Published Category", "Category", "Sub-Category", "Ingredient?", "Product?",
          "Kind", "Status", "Source", "Recipe (Made From)", "Allergen Class", "Allergen Source",
          "Type", "Unit", "Items Per Unit", "B2B Weight", "B2B Price (US$)", "B2C Weight", "B2C Price (US$)",
          "Pieces / Pack", "Notes"],
         [
             # $-string prices; recipe-library flag
             ["SAC-001", "Test Sauce", "Sauce & Marinade", "Sauce", "Cold sauce / dip", "Yes", "Yes", "Production",
              "Active", "Recipe Library", "Mayo, garlic", "Eggs, Mustard", "Eggs - mayo; Mustard - mayo",
              "KG", "KG", None, 1, "$5.00", 0.3, 3, None, "B2C bottles", "overflow-note"],
             # unpriced row
             ["SAC-002", "Test Plain Sauce", "Sauce & Marinade", "Sauce", "Cooking sauce", "Yes", "Yes", "Production",
              "Active", "Recipe Library", "Tomato, oil", "None", "-", "KG", "KG",
              None, None, None, None, None, None, None],
         ])
    fill("3. Dressing",
         ["ID", "Production Item", "Published Category", "Category", "Sub-Category", "Ingredient?", "Product?",
          "Kind", "Status", "Source", "Recipe (Made From)", "Allergen Class", "Allergen Source",
          "Type", "Usage Unit", "Items Per Unit", "B2B Weight", "B2B Price (US$)", "B2C Weight", "B2C Price (US$)",
          "Pieces / Pack", "Notes"],
         [["DRS-001", "Test Dressing", "Dressing", "Dressings", "Salad dressing", "Yes", "Yes", "Production",
           "Active", "Recipe Library", "Oil, lemon", "None", "—", "KG", "KG", None, None, None, None, None, None, None]])
    fill("4. Meals",
         ["ID", "Production Item", "Published Category", "Category", "Sub-Category", "Ingredient?", "Product?",
          "Kind", "Status", "Source", "Recipe (Made From)", "Allergen Class", "Allergen Source",
          "Type", "Unit", "Items Per Unit", "B2B Weight", "B2B Price (US$)", "B2C Weight", "B2C Price (US$)", "Notes"],
         [
             # PRD-016 pattern: composition in the Allergen Class column
             ["PRD-001", "Test Wings", "Meal", "Meat & Egg", "Poultry", "Yes", "Yes", "Production",
              "Active", "Recipe Library", "None", "Wings, garlic, spices", "-", "PIECE", "PIECE",
              None, None, None, None, None, None],
             # blank role flags, priced
             ["PRD-002", "Test Onions", "Meal", "Vegetables", "Onions & Garlic", None, None, "Production",
              "Active", "Recipe Library", "Onion, butter", "Milk", "Milk - butter", "BAG", "KG",
              None, 1, 12, 0.2, 3, "portions"],
             # name collision with a later product row
             ["PRD-003", "Test Fajita", "Meal", "Meat & Egg", "Beef", "Yes", "Yes", "Production",
              "Active", "Recipe Library", "Beef, pepper", "None", "-", "BAG", "KG",
              None, None, None, None, None, None],
         ])
    fill("5. Products",
         ["ID", "Production Item", "Published Category", "Category", "Sub-Category", "Ingredient?",
          "Kind", "Status", "Recipe (Made From)", "Allergen Class", "Allergen Source",
          "Type", "Unit", "Items Per Unit", "B2B Weight", "B2B Price (US$)", "B2C Weight", "B2C Price (US$)", "Notes"],
         [
             # composition sitting in Allergen Source; Suplier typo; slug collision
             ["RSL-001", "Test Fajita", "Frozen", "Meat & Egg", "Beef", "Yes", "Suplier", "Active",
              "Milk - cheese; Eggs", "Milk, Eggs", "Beef, cheese, pepper", None, None, None, None, None, None, None, None],
             # RSL-045 pattern: Recipe Library in the composition column
             ["RSL-002", "Test Patty", "Frozen", "Meat & Egg", "Beef", "Yes", "Suplier", "Active",
              "Recipe Library", "Minced beef, fat, onion", None, None, None, None, None, None, None, None, None],
             # RSL-055 pattern: ingredients-shaped row, shifted one column
             ["RSL-003", "Test Baguette", None, "Bread", "Wheat / gluten", "Yes", "No", "Supplier", "Active",
              "Cereals/Gluten", "Wheat flour, water", "Cereals/Gluten - wheat flour", "PIECE", None, None, None, None, None, None],
             # Ingredient?=No beverage
             ["RSL-004", "Test Cola", "Beverage", "Beverage", "Juice & Soft Drink", "No", "Suplier", "Active",
              "-", "None", "Water, sugar", None, None, None, None, None, None, None, None],
         ])
    fill("6.Packaging",
         ["ID", "Item", "Category", "Sub-Category", "Kind", "Status", "Composition (Made From)",
          "Type", "Unit", "Items per Unit", "Waste (%)"],
         [["PKG-001", "Test Bags", "Packaging & disposables", "Bags", "Supplier", "Active",
           "Polyethylene", "BAG", "PIECE", 100, 2]])
    return wb


def by_ref(doc_rows, ref):
    return next(r for r in doc_rows if r["source_ref"] == ref)


def self_test():
    platform, catalogue, repairs = convert(build_synthetic_workbook())
    ings, items = platform["ingredients"], catalogue["items"]

    assert len(ings) == 4 and len(items) == 10, (len(ings), len(items))
    assert {r["kind"] for r in ings} == {"supplier"}

    flour = by_ref(ings, "ING-001")
    assert flour["default_unit_code"] == "kg" and flour["purchase_unit_code"] == "kg"
    assert flour["allergens"][0]["allergen_code"] == "gluten"
    assert flour["allergens"][0]["evidence"] == "Cereals/Gluten - milled wheat"
    assert flour["items_per_unit"] == 1 and flour["composition"] == "Milled wheat"

    vinegar = by_ref(ings, "ING-002")
    assert vinegar["status"] == "inactive" and "status_blank_inactive" in vinegar["flags"]
    assert vinegar["default_unit_code"] == "kg" and "unit_defaulted_kg" in vinegar["flags"]
    assert vinegar["allergens"][0] == {
        "allergen_code": "sulphites", "containment": "may_contain",
        "verification_status": "requires_supplier_confirmation",
        "evidence": "Sulphites~ - vinegar"}

    coconut = by_ref(ings, "ING-003")
    assert coconut["allergens"][0]["allergen_code"] == "tree_nut"
    assert coconut["allergens"][0]["market_scope"] == "us_only"
    assert coconut["purchase_unit_code"] == "pack" and coconut["default_unit_code"] == "piece"

    # Packaging is an ingredient filed under its own branch — that category *is* the
    # discriminator, so it is the one thing here worth asserting outright.
    pkg = by_ref(ings, "PKG-001")
    assert pkg["allergens"] == [] and pkg["category_code"] == "packaging-disposables"
    assert pkg["subcategory_code"] == "packaging-disposables-bags"
    assert pkg["default_unit_code"] == "piece" and pkg["purchase_unit_code"] == "bag"
    assert pkg["items_per_unit"] == 100 and pkg["waste_percent"] == 2
    # The sheet carries no cost column; the curated overlay is what supplies one, and a row the
    # overlay does not mention keeps exactly what the sheet said.
    assert "purchase_price_amount" not in pkg

    priced = convert_packaging(
        build_synthetic_workbook()["6.Packaging"], Taxonomy(), SlugBook(),
        {"PKG-001": {"purchase_price_amount": 6.5, "purchase_price_currency": "USD"}},
    )[0]
    assert priced["purchase_price_amount"] == 6.5 and priced["purchase_price_currency"] == "USD"

    sauce = by_ref(items, "SAC-001")
    assert sauce["prices"]["b2b"] == {"weight_kg": 1, "price_minor": 500, "currency": "USD"}
    assert sauce["prices"]["b2c"] == {"weight_kg": 0.3, "price_minor": 300, "currency": "USD"}
    assert "recipe_library_unlinked" in sauce["flags"] and sauce["notes"] == "B2C bottles | overflow-note"
    assert sauce["published_category_code"] == "sauce"
    assert by_ref(items, "SAC-002")["prices"] == {}
    assert by_ref(items, "DRS-001")["sheet_item_type"] == "dressing"

    wings = by_ref(items, "PRD-001")
    assert wings["composition"] == "Wings, garlic, spices" and wings["allergens"] == []
    assert "columns_swapped_repaired" in wings["flags"]

    onions = by_ref(items, "PRD-002")
    assert onions["is_product"] and not onions["is_ingredient"]
    assert "source_blank_role_flags" in onions["flags"]
    assert onions["prices"]["b2b"]["price_minor"] == 1200

    meal_fajita, product_fajita = by_ref(items, "PRD-003"), by_ref(items, "RSL-001")
    assert meal_fajita["slug"] == "test-fajita" and product_fajita["slug"] == "test-fajita-product"
    assert product_fajita["composition"] == "Beef, cheese, pepper"
    assert {m["allergen_code"] for m in product_fajita["allergens"]} == {"milk", "egg"}
    assert product_fajita["production_mode"] == "supplier"
    assert product_fajita["usage_unit_code"] == "piece" and "unit_defaulted_piece" in product_fajita["flags"]

    patty = by_ref(items, "RSL-002")
    assert patty["source"] == "Recipe Library" and patty["composition"] == "Minced beef, fat, onion"
    assert "recipe_library_unlinked" in patty["flags"]

    baguette = by_ref(items, "RSL-003")
    assert "row_shift_repaired" in baguette["flags"]
    assert baguette["status"] == "active" and baguette["composition"] == "Wheat flour, water"
    assert baguette["allergens"][0]["allergen_code"] == "gluten"
    assert baguette["usage_unit_code"] == "piece" and "published_category_missing" in baguette["flags"]

    cola = by_ref(items, "RSL-004")
    assert not cola["is_ingredient"] and cola["composition"] == "Water, sugar"

    cats = {c["code"]: c for c in platform["categories"]}
    assert "vegetables" not in cats and cats["vegetable"]["name_en"] == "Vegetable", "Vegetables folds onto Vegetable"
    assert cats["meat-egg-poultry"]["parent_code"] == "meat-egg"
    assert repairs, "repairs are reported"
    print(f"self-test OK ({len(ings)} ingredients, {len(items)} items, {len(repairs)} repairs exercised)")



# ---------------------------------------------------------------- recipes mode

RECIPE_WORKBOOKS = {
    "sauce": "Actual Data_Recipes (Sauce).xlsx",
    "dressing": "Actual Data_Recipes (Dressing).xlsx",
    "meal": "Actual Data_Recipes (Meal).xlsx",
}


def parse_yield(raw) -> dict:
    """The Quantity Produced cell in its three shapes: a bare number (kg or
    pieces - the cost labels disambiguate downstream), "130 (7 kg)" (pieces
    with the mass in parentheses), or empty."""
    text = clean(raw)
    out = {"raw": text or None, "label": None, "shape": "absent",
           "quantity": None, "unit": None, "piece_count": None}
    if not text:
        return out
    m = re.match(r"^(\d+)\s*\(\s*([\d.]+)\s*kg\s*\)$", text, re.IGNORECASE)
    if m:
        out.update(shape="pieces_with_mass", piece_count=int(m.group(1)),
                   quantity=m.group(2), unit="kg")
        return out
    if re.match(r"^[\d.]+$", text):
        out.update(shape="bare_number", quantity=text, unit="kg")
        return out
    out["shape"] = "unparsed"
    return out


def convert_recipe_sheet(ws, family: str, index: int) -> dict:
    rows = []
    for r in ws.iter_rows(values_only=True):
        r = [clean(c) for c in r]
        while r and r[-1] == "":
            r.pop()
        if r:
            rows.append(r)

    sheet = {
        "sheet_index": index, "family": family, "tab": ws.title,
        "designation": "", "kind": None,
        "yield": parse_yield(None),
        "totals": {"input_quantity": None, "input_total": None},
        "lines": [], "cost_labels": [], "waste_percent": None,
        "findings": [],
    }
    in_lines = False
    for row_number, r in enumerate(rows, start=1):
        head = r[0]
        low = head.lower()
        if head == "Designation" and len(r) > 1 and r[1] == "U.":
            in_lines = True
            continue
        if low.startswith("total:"):
            in_lines = False
            if len(r) > 2:
                sheet["totals"]["input_quantity"] = clean(r[2]) or None
            if len(r) > 4:
                sheet["totals"]["input_total"] = clean(r[4]) or None
            continue
        if in_lines:
            if len(r) >= 2 and r[0]:
                sheet["lines"].append({
                    "row": row_number,
                    "designation": r[0],
                    "unit": r[1] if len(r) > 1 and r[1] else None,
                    "quantity": r[2] if len(r) > 2 and r[2] else None,
                    "unit_price": r[3] if len(r) > 3 and r[3] else None,
                    "line_total": r[4] if len(r) > 4 and r[4] else None,
                    "comment": r[5] if len(r) > 5 and r[5] else None,
                })
            continue
        if low == "designation" and len(r) > 1 and not sheet["designation"]:
            sheet["designation"] = r[1]
        elif low == "kind" and len(r) > 1:
            sheet["kind"] = r[1]
        elif low.startswith("quantity produced") and len(r) > 1:
            sheet["yield"] = parse_yield(r[1])
        elif ("production cost" in low or "waste" in low) and len(r) > 1:
            sheet["cost_labels"].append({"label": head, "amount": clean(r[1]) or None})
            if "waste" in low:
                m = re.search(r"(\d+(?:\.\d+)?)\s*%", head)
                if m:
                    sheet["waste_percent"] = m.group(1)

    if not sheet["designation"]:
        sheet["designation"] = ws.title
        sheet["findings"].append({"code": "designation_absent",
                                  "detail": f"{family}/{ws.title}: no Designation cell; the tab name stands in."})
    if sheet["yield"]["shape"] in {"absent", "unparsed"}:
        code = "yield_unstated" if sheet["yield"]["shape"] == "absent" else "yield_unparsed"
        sheet["findings"].append({"code": code,
                                  "detail": f"{family}/{ws.title}: Quantity Produced is {sheet['yield']['raw'] or 'empty'}."})
    return sheet


def convert_recipes(directory: str) -> dict:
    import openpyxl

    sheets, findings = [], []
    index = 0
    for family, filename in RECIPE_WORKBOOKS.items():
        wb = openpyxl.load_workbook(str(Path(directory) / filename), data_only=True)
        for tab in wb.sheetnames:
            index += 1
            sheets.append(convert_recipe_sheet(wb[tab], family, index))

    # Repairs the sheets need before the writer sees them, each reported.
    # Order matters: dedupe first, so a sheet pasted twice under two tabs is
    # recognised as the same sheet, not mistaken for a name collision.
    #
    # 1. The same sheet pasted into two workbooks (Honey Mustard, equal
    #    figures) imports once; a second *different* formulation would stay -
    #    the writer models that as a second draft version, deliberately.
    unique, by_key = [], {}
    for sheet in sheets:
        key = (sheet["designation"].lower(), sheet["totals"]["input_total"], sheet["yield"]["raw"])
        if key in by_key:
            findings.append({"code": "duplicate_sheet_skipped",
                             "detail": f'{sheet["family"]}/{sheet["tab"]}: identical to {by_key[key]}; imported once.'})
            continue
        by_key[key] = f'{sheet["family"]}/{sheet["tab"]}'
        unique.append(sheet)

    # 2. A designation another sheet already claims while the tab and the
    #    figures both disagree - the Ranch sheet says "Caesar Sauce" - is a
    #    copy-paste slip; the tab wins.
    seen = {}
    for sheet in unique:
        key = sheet["designation"].lower()
        if key in seen and sheet["tab"].lower() != key:
            findings.append({"code": "designation_collision_repaired",
                             "detail": f'{sheet["family"]}/{sheet["tab"]}: Designation says "{sheet["designation"]}" - already sheet {seen[key]}. The tab name stands in.'})
            sheet["designation"] = sheet["tab"]
        seen.setdefault(sheet["designation"].lower(), f'{sheet["family"]}/{sheet["tab"]}')
    for i, sheet in enumerate(unique, start=1):
        sheet["sheet_index"] = i

    return {
        "source_workbooks": list(RECIPE_WORKBOOKS.values()),
        "source_system": "healthy360_workbook_v6",
        "sheets": unique,
        "findings": findings,
    }


def recipes_self_test():
    import openpyxl
    import tempfile

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def sheet(title, designation, kind, qty, lines, costs):
        ws = wb.create_sheet(title)
        ws.append(["Technical Sheet"])
        ws.append(["Description"])
        ws.append(["Designation", designation, "", "", "Insert Photo Here"])
        ws.append(["Kind", kind])
        ws.append(["Quantity Produced", qty])
        ws.append(["Production"])
        ws.append(["Raw Materiel"])
        ws.append(["Designation", "U.", "Q.", "U.P.", "T", "Comments"])
        total_q = total_t = 0
        for name, q, up in lines:
            ws.append([name, "kg", q, up, q * up])
            total_q += q
            total_t += q * up
        ws.append(["Total:", "", total_q, "", total_t])
        ws.append(["Cost"])
        for label, amount in costs:
            ws.append([label, amount])

    sheet("Test Sauce", "Test Sauce", "Production", 1.5,
          [("Mayo", 1, 3.5), ("Salt", 0.01, 0.3)],
          [("Total Production Cost", 3.503), ("1 kg Production Cost", 2.335), ("3% Waste Coeffecient", 2.405)])
    sheet("Piece Prep", "Piece Prep", "Preparation", "130 (7 kg)",
          [("Chicken", 7, 4)],
          [("Total Production Cost", 28), ("1 Piece Production Cost", 0.215), ("3% Waste Coeffecient", 0.222),
           ("1 Kg Production Cost", 4.0), ("3% Waste Coeffecient", 4.12)])
    sheet("Ranch Sheet", "Test Sauce", "Sauce", 2,
          [("Mayo", 2, 3.5)],
          [("Total Production Cost", 7), ("Add Waste Coefficient 3%", 3.605)])
    sheet("Dup Sauce", "Test Sauce", "Production", 1.5,
          [("Mayo", 1, 3.5), ("Salt", 0.01, 0.3)],
          [("Total Production Cost", 3.503)])

    with tempfile.TemporaryDirectory() as d:
        wb.save(str(Path(d) / RECIPE_WORKBOOKS["sauce"]))
        for family in ("dressing", "meal"):
            small = openpyxl.Workbook()
            small.remove(small.active)
            ws = small.create_sheet("Filler")
            ws.append(["Technical Sheet"])
            ws.append(["Designation", f"Filler {family}"])
            ws.append(["Kind", "Production"])
            ws.append(["Quantity Produced", 1])
            ws.append(["Designation", "U.", "Q.", "U.P.", "T"])
            ws.append(["Salt", "kg", 1, 0.3, 0.3])
            ws.append(["Total:", "", 1, "", 0.3])
            ws.append(["Total Production Cost", 0.3])
            small.save(str(Path(d) / RECIPE_WORKBOOKS[family]))
        doc = convert_recipes(d)

    sheets = {s["tab"]: s for s in doc["sheets"]}
    assert len(doc["sheets"]) == 5, len(doc["sheets"])  # Dup Sauce dropped
    assert "Dup Sauce" not in sheets
    assert any(f["code"] == "duplicate_sheet_skipped" for f in doc["findings"])
    assert sheets["Ranch Sheet"]["designation"] == "Ranch Sheet"
    assert any(f["code"] == "designation_collision_repaired" for f in doc["findings"])
    test = sheets["Test Sauce"]
    assert test["yield"] == {"raw": "1.5", "label": None, "shape": "bare_number",
                             "quantity": "1.5", "unit": "kg", "piece_count": None}
    assert test["waste_percent"] == "3"
    assert test["totals"]["input_quantity"] == "1.01"
    assert test["lines"][0]["designation"] == "Mayo" and test["lines"][0]["unit_price"] == "3.5"
    piece = sheets["Piece Prep"]
    assert piece["yield"]["piece_count"] == 130 and piece["yield"]["quantity"] == "7"
    assert len(piece["cost_labels"]) == 5
    assert [s["sheet_index"] for s in doc["sheets"]] == [1, 2, 3, 4, 5]
    print(f"recipes self-test OK ({len(doc['sheets'])} sheets, {len(doc['findings'])} findings exercised)")

# ---------------------------------------------------------------- entrypoint

def main():
    self_test()
    recipes_self_test()
    if "--self-test" in sys.argv:
        return

    if "--recipes" in sys.argv:
        directory = sys.argv[sys.argv.index("--recipes") + 1]
        out = Path(sys.argv[sys.argv.index("--recipes-out") + 1]) if "--recipes-out" in sys.argv \
            else Path(directory) / "v6-recipes.json"
        doc = convert_recipes(directory)
        # Confidential: formulations and unit costs. Written beside the source
        # workbooks (or --recipes-out), NEVER into the repository.
        out.write_text(json.dumps(doc, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        lines = sum(len(s["lines"]) for s in doc["sheets"])
        print(f"wrote {out} ({len(doc['sheets'])} recipe sheets, {lines} lines) - PRIVATE, do not commit")
        for finding in doc["findings"]:
            print(f"  - [{finding['code']}] {finding['detail']}")
        return

    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        print("usage: convert-v6-workbook.py <path-to-xlsx> | --recipes <dir> [--recipes-out <file>] | --self-test", file=sys.stderr)
        sys.exit(2)

    import openpyxl

    workbook = openpyxl.load_workbook(args[0], data_only=True)

    # The curated price/capacity overlay for packaging, keyed by source_ref. Sheet 6 carries no
    # cost column, so without this every box regenerates unpriced — and a technical sheet
    # withholds its total entirely while one packaging line has none.
    overlay = {}
    if PACKAGING_OVERLAY.exists():
        overlay = {
            row["source_ref"]: row
            for row in json.loads(PACKAGING_OVERLAY.read_text(encoding="utf-8"))["items"]
        }
    else:
        print(f"warning: {PACKAGING_OVERLAY} is missing; packaging will regenerate unpriced",
              file=sys.stderr)

    platform_doc, catalogue_doc, repairs = convert(workbook, overlay)

    ings, items = platform_doc["ingredients"], catalogue_doc["items"]
    assert len(ings) == 337, f"expected 306 ING + 31 PKG rows, got {len(ings)}"
    assert len(items) == 155, f"expected 43+19+38+55 items, got {len(items)}"
    assert len({i["source_ref"] for i in ings}) == len(ings)
    assert len({i["slug"] for i in ings}) == len(ings)
    assert len({i["source_ref"] for i in items}) == len(items)
    assert len({i["slug"] for i in items}) == len(items)

    CATALOGUE_OUT.parent.mkdir(parents=True, exist_ok=True)
    PLATFORM_OUT.write_text(json.dumps(platform_doc, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    CATALOGUE_OUT.write_text(json.dumps(catalogue_doc, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    unmapped = [(i["source_ref"], f) for i in ings + items for f in i.get("flags", []) if f.startswith("allergen_unmapped")]
    print(f"wrote {PLATFORM_OUT.relative_to(REPO_ROOT)} ({len(ings)} ingredients, {len(platform_doc['categories'])} categories)")
    print(f"wrote {CATALOGUE_OUT.relative_to(REPO_ROOT)} ({len(items)} items)")
    print(f"{len(repairs)} row repairs applied:")
    for line in repairs:
        print(f"  - {line}")
    if unmapped:
        print(f"{len(unmapped)} unmapped allergen wordings (dropped, flagged):")
        for ref, flag in unmapped:
            print(f"  - {ref}: {flag}")


if __name__ == "__main__":
    main()
